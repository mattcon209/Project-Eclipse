#!/usr/bin/env python3
"""Shared Word styling for project specification documents."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
STEEL = RGBColor(0x2C, 0x3E, 0x5A)
ACCENT = RGBColor(0x1A, 0x6B, 0x5C)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
SLATE = RGBColor(0x4A, 0x55, 0x68)
MUTED = RGBColor(0x6B, 0x73, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_BLACK = RGBColor(0x1A, 0x1A, 0x1A)
TEAL_LT = RGBColor(0x7D, 0xD3, 0xC0)
HEADER_BG = "1B2A4A"
ROW_ALT = "F4F6F8"
WARN_BG = "FBF4E4"
WARN_BORDER = "C9A227"
LIGHT_TEAL = "E8F4F1"
TEAL_BORDER = "1A6B5C"
GOLD_ROW = "FBF4E4"
NAVY_MID = "243656"
ICE = "F7FAFC"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=NEAR_BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, color="D0D5DD", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_inches=7.0):
    table.autofit = False
    table.allow_autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(width_inches * 1440)))
    tblW.set(qn("w:type"), "dxa")
    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")


def set_col_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            cell = row.cells[i]
            cell.width = Inches(w)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(w * 1440)))
            tcW.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color="1A6B5C", sz="18"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def set_paragraph_spacing(p, before=0, after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_fld(paragraph, instr, size=8, color=MUTED):
    r1 = paragraph.add_run()
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc1)
    r2 = paragraph.add_run()
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    r2._r.append(instrText)
    set_run_font(r2, size=size, color=color)
    r3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "end")
    r3._r.append(fc2)


def setup_page(doc, header_kicker, header_sub, footer_left):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(1.05)
    section.bottom_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 2)
    run = p.add_run(header_kicker)
    set_run_font(run, size=8, bold=True, color=ACCENT)
    p2 = header.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p2, 0, 4)
    run = p2.add_run(header_sub)
    set_run_font(run, size=8, italic=True, color=MUTED)
    pPr = p2._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "1B2A4A")
    pBdr.append(bottom)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, 4, 0)
    run = p.add_run(footer_left)
    set_run_font(run, size=8, color=MUTED, italic=True)
    run2 = p.add_run("\t")
    set_run_font(run2, size=8, color=MUTED)
    pPr = p._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    pBdr = OxmlElement("w:pBdr")
    pPr.append(pBdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "8")
    top.set(qn("w:color"), "1A6B5C")
    pBdr.append(top)
    run3 = p.add_run("Page ")
    set_run_font(run3, size=8, color=MUTED)
    add_fld(p, " PAGE ")
    run4 = p.add_run(" of ")
    set_run_font(run4, size=8, color=MUTED)
    add_fld(p, " NUMPAGES ")

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = NEAR_BLACK
    return section


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        set_paragraph_spacing(p, 16, 8)
        run = p.add_run(text.upper())
        set_run_font(run, size=13, bold=True, color=NAVY)
        add_bottom_border(p, "1A6B5C", "16")
    elif level == 2:
        set_paragraph_spacing(p, 12, 6)
        run = p.add_run(text)
        set_run_font(run, size=12, bold=True, color=STEEL)
        add_bottom_border(p, "D0D5DD", "8")
    else:
        set_paragraph_spacing(p, 8, 4)
        run = p.add_run(text)
        set_run_font(run, size=11, bold=True, color=ACCENT)
    return p


def body(doc, text, size=11, italic=False, color=NEAR_BLACK, after=8, before=0, bold=False):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before, after, 1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, color=color, bold=bold)
    return p


def mixed(doc, parts, after=8, before=0, size=11):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before, after, 1.15)
    for text, bold, italic, color in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def bullet(doc, text, bold_lead=None, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.35 + level * 0.25)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, size=11, bold=True, color=NEAR_BLACK)
        r = p.add_run(text)
        set_run_font(r, size=11, color=NEAR_BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=NEAR_BLACK)
    return p


def spacer(doc, pts=6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, 0, pts, 1)
    return p


def fill_header_row(row, headers, bg=HEADER_BG):
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, bg)
        set_cell_borders(cell, bg, "4")
        set_cell_margins(cell, 70, 70, 90, 90)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, 0, 0)
        run = p.add_run(h)
        set_run_font(run, size=9.5, bold=True, color=WHITE)


def fill_cell(cell, text, *, bg="FFFFFF", bold=False, color=NEAR_BLACK, size=9.5, border="E4E7EC"):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, bg)
    set_cell_borders(cell, border, "4")
    set_cell_margins(cell, 60, 60, 90, 90)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, 0, 0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def simple_table(doc, headers, rows, widths, first_col_bold=True, alt_bg=ROW_ALT):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 7.0)
    set_col_widths(table, widths)
    fill_header_row(table.rows[0], headers)
    for idx, row_data in enumerate(rows):
        bg = alt_bg if idx % 2 == 1 else "FFFFFF"
        for ci, val in enumerate(row_data):
            fill_cell(
                table.rows[idx + 1].cells[ci],
                val,
                bg=bg,
                bold=(first_col_bold and ci == 0),
                color=STEEL if (first_col_bold and ci == 0) else NEAR_BLACK,
            )
    return table


def kv_table(doc, rows, col_widths=(2.2, 4.8), header=("Property", "Value")):
    return simple_table(doc, list(header), rows, list(col_widths))


def banner_cell(doc, lines, bg=HEADER_BG, border="1B2A4A"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 7.0)
    set_col_widths(table, [7.0])
    cell = table.rows[0].cells[0]
    cell.text = ""
    set_cell_shading(cell, bg)
    set_cell_borders(cell, border, "4")
    set_cell_margins(cell, 180, 180, 180, 180)
    first = True
    for text, size, bold, italic, color, after, align in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = align
        set_paragraph_spacing(p, 0, after)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return table


def callout(doc, title, text, bg=WARN_BG, title_color=GOLD, border=WARN_BORDER):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 7.0)
    set_col_widths(table, [7.0])
    cell = table.rows[0].cells[0]
    cell.text = ""
    set_cell_shading(cell, bg)
    set_cell_borders(cell, border, "12")
    set_cell_margins(cell, 120, 120, 140, 140)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, 0, 4)
    run = p.add_run(title)
    set_run_font(run, size=10, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, 0, 0, 1.15)
    run = p2.add_run(text)
    set_run_font(run, size=10, color=NEAR_BLACK)
    return table


def state_table(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 7.0)
    set_col_widths(table, [1.2, 5.8])
    fill_header_row(table.rows[0], ["State", "Item"])
    for idx, (state, item) in enumerate(rows):
        bg = LIGHT_TEAL if state == "Locked" else GOLD_ROW
        fill_cell(
            table.rows[idx + 1].cells[0],
            state,
            bg=bg,
            bold=True,
            color=ACCENT if state == "Locked" else GOLD,
        )
        fill_cell(table.rows[idx + 1].cells[1], item, bg=bg)
    return table


def new_doc():
    return Document()
